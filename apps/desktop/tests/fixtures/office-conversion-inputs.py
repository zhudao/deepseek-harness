"""Create one-page Office inputs with the packaged authoring runtime."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

root = Path(sys.argv[1])
document = Document()
document.add_paragraph("Desktop Office conversion 42")
document.save(root / "input.docx")
workbook = Workbook()
workbook.active.append(["Desktop Office conversion", 42])
workbook.save(root / "input.xlsx")
workbook.close()
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Desktop Office conversion 42"
presentation.save(root / "input.pptx")
