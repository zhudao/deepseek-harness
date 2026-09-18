"""Exercise the relocated Office payload with its own isolated interpreter."""

import importlib.metadata
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile

import numpy
import pandas
from docx import Document
from docx.shared import Inches as DocxInches
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def main():
    """Check installed versions and read back editable Office documents."""
    assert sys.version_info[:3] == tuple(map(int, sys.argv[2].split("."))), sys.version
    versions = json.loads(sys.argv[1])
    # Only pip belongs to the interpreter baseline; all other distributions must be declared.
    # Each target's native release-host smoke must confirm this baseline.
    expected_names = {re.sub(r"[-_.]+", "-", name).lower() for name in versions} | {"pip"}
    installed_names = {
        re.sub(r"[-_.]+", "-", distribution.metadata["Name"]).lower()
        for distribution in importlib.metadata.distributions()
    }
    assert installed_names == expected_names, {"unexpected": sorted(installed_names - expected_names), "missing": sorted(expected_names - installed_names)}
    for name, expected in versions.items():
        actual = importlib.metadata.version(name)
        assert actual == expected, (name, actual, expected)
    assert numpy.arange(4).sum() == 6
    assert pandas.DataFrame({"n": [1, 2]}).n.sum() == 3

    with tempfile.TemporaryDirectory(prefix="dsh-office-smoke-") as directory:
        root = Path(directory)
        image = root / "chart.png"
        Image.new("RGB", (80, 40), "#2878bc").save(image)
        with Image.open(image) as restored:
            assert restored.size == (80, 40)

        document = Document()
        document.add_heading("Office 文档", 0)
        document.add_paragraph("Editable text")
        document.add_table(rows=2, cols=2).cell(1, 1).text = "42"
        document.add_picture(str(image), width=DocxInches(1))
        document.save(root / "document.docx")
        reopened_document = Document(root / "document.docx")
        assert reopened_document.tables[0].cell(1, 1).text == "42"
        assert reopened_document.paragraphs[0].text == "Office 文档"

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Office 演示"
        slide.shapes.add_picture(str(image), Inches(1), Inches(2))
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("Values", [2, 4])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(2), Inches(4), Inches(3), chart_data)
        presentation.save(root / "presentation.pptx")
        reopened_presentation = Presentation(root / "presentation.pptx")
        assert len(reopened_presentation.slides) == 1
        chart = next(shape.chart for shape in reopened_presentation.slides[0].shapes if shape.has_chart)
        assert list(chart.series[0].values) == [2.0, 4.0]

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Value", "Formula"])
        sheet.append([42, "=A2*2"])
        sheet["A1"].font = Font(bold=True)
        workbook.save(root / "workbook.xlsx")
        reopened_workbook = load_workbook(root / "workbook.xlsx")
        try:
            assert reopened_workbook.active["A2"].value == 42
            assert reopened_workbook.active["B2"].value == "=A2*2"
            assert reopened_workbook.active["A1"].font.bold
        finally:
            reopened_workbook.close()
        assert pandas.read_excel(root / "workbook.xlsx")["Value"].iloc[0] == 42
        for file, arguments in [
            ("document.docx", ["--contains", "Office 文档"]),
            ("presentation.pptx", ["--contains", "Office 演示", "--count", "1"]),
            ("workbook.xlsx", ["--contains", "Value", "--count", "1"]),
        ]:
            checked = subprocess.run([sys.executable, "-I", "-B", sys.argv[3], str(root / file), *arguments],
                                     check=False, capture_output=True, text=True, timeout=30)
            assert checked.returncode == 0 and json.loads(checked.stdout)["verdict"] == "pass", checked.stdout + checked.stderr
    print("Office runtime versions and document round trips passed.")


if __name__ == "__main__":
    main()
