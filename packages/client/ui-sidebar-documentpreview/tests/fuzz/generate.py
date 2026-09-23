"""Generate replayable XLSX cases using independent writers and equivalent ZIP/XML encodings."""

import argparse
import base64
import copy
from datetime import datetime, time, timedelta
import io
import json
import math
from pathlib import Path
import posixpath
import random
import re
import warnings
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904, to_excel
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
import xlsxwriter

FEATURES = ['chart', 'image', 'comment', 'conditional', 'conditional-databar', 'conditional-iconset', 'conditional-expression', 'table', 'validation', 'hyperlink',
            'merge', 'freeze', 'hidden', 'header', 'epoch1904', 'styled']
ERRORS = ['#NULL!', '#DIV/0!', '#VALUE!', '#REF!', '#NAME?', '#NUM!', '#N/A',
          '#SPILL!', '#CALC!', '#GETTING_DATA']
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
MAIN_NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def fixed_zip(files, compression=zipfile.ZIP_DEFLATED, reverse=False, comment=b''):
    """Fix ZIP metadata so the same seed yields byte-identical files."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=compression) as archive:
        for name in sorted(files, reverse=reverse):
            entry = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            entry.compress_type = compression
            content = files[name]
            if name == 'docProps/core.xml':
                # Openpyxl overwrites the modified time during save_workbook.
                content = re.sub(rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)', rb'\g<1>2020-01-01T00:00:00Z\g<2>', content)
            archive.writestr(entry, content)
        archive.comment = comment
    return output.getvalue()


def unpack(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def validate(data, cells):
    """Check internal relationship targets and independently reload all oracle cells."""
    files = unpack(data)
    for name, content in files.items():
        if name.endswith('.rels'):
            folder = '' if name == '_rels/.rels' else name.split('/_rels/')[0]
            for rel in ET.fromstring(content):
                if rel.get('TargetMode') == 'External':
                    continue
                target = rel.attrib['Target']
                target = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(folder, target))
                if target not in files:
                    raise ValueError(f'Missing relationship target {name}: {target}')
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    for cell in cells:
        actual = workbook[cell['sheet']].cell(cell['r'] + 1, cell['c'] + 1).value
        expected = cell.get('formula', cell.get('source', cell.get('value')))
        if isinstance(expected, dict):
            continue
        equivalent_number = type(actual) in (int, float) and type(expected) in (int, float) and math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)
        if actual != expected and not equivalent_number:
            raise ValueError(f'Writer changed {cell}: {actual!r}')
    workbook.close()


def workbook_bytes(writer, features, values, case_seed):
    """Keep expected values independent of the preview implementation."""
    output = io.BytesIO()
    cells = []
    active = set(features)
    if writer == 'openpyxl':
        book = openpyxl.Workbook()
        book.properties.created = book.properties.modified = datetime(2020, 1, 1)
        if 'epoch1904' in active:
            book.epoch = CALENDAR_MAC_1904
        sheet = book.active
        sheet.title = '数据'
    else:
        book = xlsxwriter.Workbook(output, {'in_memory': True, 'date_1904': 'epoch1904' in active})
        book.set_properties({'created': datetime(2020, 1, 1)})
        sheet = book.add_worksheet('数据')

    for r, c, value, kind in values:
        oracle = {'sheet': '数据', 'r': r, 'c': c}
        if kind == 'formula':
            formula, cached = value
            oracle['formula'] = formula
            oracle['value'] = None if writer == 'openpyxl' else cached
            if writer == 'openpyxl':
                sheet.cell(r + 1, c + 1, formula)
            else:
                sheet.write_formula(r, c, formula, None, cached)
        elif kind == 'date':
            oracle['source'] = {'iso': value.isoformat()}
            oracle['value'] = to_excel(value)
            if writer == 'openpyxl':
                sheet.cell(r + 1, c + 1, value).number_format = 'yyyy-mm-dd hh:mm:ss'
            else:
                sheet.write_datetime(r, c, value, book.add_format({'num_format': 'yyyy-mm-dd hh:mm:ss'}))
        elif kind == 'time':
            oracle['source'] = {'iso': str(value)}
            oracle['value'] = to_excel(value)
            if writer == 'openpyxl':
                sheet.cell(r + 1, c + 1, value).number_format = '[h]:mm:ss' if isinstance(value, timedelta) else 'hh:mm:ss'
            else:
                sheet.write_datetime(r, c, value, book.add_format({'num_format': '[h]:mm:ss'}))
        elif kind == 'error':
            oracle['value'] = value
            if writer == 'openpyxl':
                cell = sheet.cell(r + 1, c + 1, value)
                cell.data_type = 'e'
            else:
                # XlsxWriter exposes Excel error caches through its formula writer.
                sheet.write_formula(r, c, '=1/0', None, value)
                oracle['formula'] = '=1/0'
        else:
            oracle['value'] = value
            if writer == 'openpyxl':
                sheet.cell(r + 1, c + 1, value)
            elif isinstance(value, str):
                sheet.write_string(r, c, value)
            else:
                sheet.write(r, c, value)
        cells.append(oracle)

    if writer == 'openpyxl':
        if 'chart' in active:
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            sheet.add_chart(chart, 'H2')
        if 'image' in active:
            sheet.add_image(Image(io.BytesIO(PNG)), 'H18')
        if 'comment' in active:
            sheet['B2'].comment = Comment('生成的批注', 'fuzz')
        if 'conditional' in active:
            sheet.conditional_formatting.add('B2:B3', ColorScaleRule(start_type='min', start_color='FF0000', end_type='max', end_color='00FF00'))
        if 'conditional-databar' in active:
            sheet.conditional_formatting.add('B2:B3', DataBarRule(start_type='min', end_type='max', color='FF336699'))
        if 'conditional-iconset' in active:
            sheet.conditional_formatting.add('B2:B3', IconSetRule('3TrafficLights1', 'percent', [0, 33, 67]))
        if 'conditional-expression' in active:
            sheet.conditional_formatting.add('B2:B3', FormulaRule(formula=['B2>20'], font=Font(bold=True)))
        if 'table' in active:
            table = Table(displayName='DataTable', ref='A1:B3')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
        if 'validation' in active:
            validation = DataValidation(type='list', formula1='"A,B,C"')
            sheet.add_data_validation(validation)
            validation.add('F1:F3')
        if 'hyperlink' in active:
            sheet['A2'].hyperlink = 'https://example.com/?a=1&b=2'
        if 'merge' in active:
            sheet.merge_cells('H30:J30')
        if 'freeze' in active:
            sheet.freeze_panes = 'B2'
        if 'hidden' in active:
            sheet.row_dimensions[3].hidden = True
            sheet.column_dimensions['B'].hidden = True
            book.create_sheet('隐藏页').sheet_state = 'hidden'
        if 'header' in active:
            sheet.oddHeader.center.text = 'Fuzz &P'
        if 'styled' in active:
            sheet['B2'].font = Font(bold=True, color='FF112233')
            sheet['B2'].fill = PatternFill('solid', fgColor='FFFFFF00')
        book.save(output)
    else:
        if 'chart' in active:
            chart = book.add_chart({'type': 'column'})
            chart.add_series({'values': ['数据', 1, 1, 2, 1]})
            sheet.insert_chart('H2', chart)
        if 'image' in active:
            sheet.insert_image('H18', 'pixel.png', {'image_data': io.BytesIO(PNG)})
        if 'comment' in active:
            sheet.write_comment('B2', '生成的批注')
        if 'conditional' in active:
            sheet.conditional_format('B2:B3', {'type': '2_color_scale'})
        if 'conditional-databar' in active:
            sheet.conditional_format('B2:B3', {'type': 'data_bar', 'data_bar_2010': True})
        if 'conditional-iconset' in active:
            sheet.conditional_format('B2:B3', {'type': 'icon_set', 'icon_style': '3_traffic_lights'})
        if 'conditional-expression' in active:
            sheet.conditional_format('B2:B3', {'type': 'formula', 'criteria': '=B2>20', 'format': book.add_format({'bold': True})})
        if 'textbox' in active:
            sheet.insert_textbox('H32', 'Text box <chart/>')
        if 'table' in active:
            sheet.add_table('A1:B3', {'columns': [{'header': 'Item'}, {'header': 'Amount'}]})
        if 'validation' in active:
            sheet.data_validation('F1:F3', {'validate': 'list', 'source': ['A', 'B', 'C']})
        if 'hyperlink' in active:
            sheet.write_url('A2', 'https://example.com/?a=1&b=2', string='A')
        if 'merge' in active:
            sheet.merge_range('H30:J30', '', book.add_format())
        if 'freeze' in active:
            sheet.freeze_panes(1, 1)
        if 'hidden' in active:
            sheet.set_row(2, None, None, {'hidden': True})
            sheet.set_column('B:B', None, None, {'hidden': True})
            book.add_worksheet('隐藏页').hide()
        if 'header' in active:
            sheet.set_header('&CFuzz &P')
        if 'styled' in active:
            sheet.write('B2', 42, book.add_format({'bold': True, 'font_color': '#112233', 'bg_color': '#FFFF00'}))
        book.close()
    features = [key for feature, key in [('chart', 'charts'), ('image', 'images'), ('textbox', 'shapes')] if feature in active]
    if any(feature.startswith('conditional') for feature in active):
        features.append('conditionalFormatting')
    return fixed_zip(unpack(output.getvalue())), cells, features


def mutate(data, mutation):
    """Change encodings or OPC names without changing worksheet cell semantics."""
    files = unpack(data)
    if mutation == 'zip-store':
        return fixed_zip(files, zipfile.ZIP_STORED)
    if mutation == 'zip-reverse-comment':
        return fixed_zip(files, reverse=True, comment=b'Excel fuzz archive comment')
    if mutation == 'sheet-prefix':
        path = 'xl/worksheets/sheet1.xml'
        xml = files[path].decode()
        xml = xml.replace(f'xmlns="{MAIN_NS}"', f'xmlns:s="{MAIN_NS}"')
        files[path] = re.sub(r'<(/?)([A-Za-z_][\w.-]*)(?=[\s/>])', r'<\1s:\2', xml).encode()
    if mutation == 'sheet-filename':
        old, new = 'xl/worksheets/sheet1.xml', 'xl/worksheets/budget.xml'
        files[new] = files.pop(old)
        rel = 'xl/worksheets/_rels/sheet1.xml.rels'
        if rel in files:
            files['xl/worksheets/_rels/budget.xml.rels'] = files.pop(rel)
        for path in ['xl/_rels/workbook.xml.rels', '[Content_Types].xml']:
            files[path] = files[path].replace(b'worksheets/sheet1.xml', b'worksheets/budget.xml')
    if mutation == 'relationship-dot':
        path = 'xl/_rels/workbook.xml.rels'
        files[path] = files[path].replace(b'Target="worksheets/', b'Target="./worksheets/').replace(b'Target="/xl/worksheets/', b'Target="/xl/./worksheets/')
    if mutation == 'relationship-absolute':
        path = 'xl/_rels/workbook.xml.rels'
        files[path] = files[path].replace(b'Target="worksheets/', b'Target="/xl/worksheets/')
    if mutation in ['drawing-default', 'drawing-alternate']:
        for path in list(files):
            if re.fullmatch(r'xl/drawings/[^/]+\.xml', path):
                xml = files[path].decode()
                match = re.search(r'xmlns(?::([\w]+))?="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"', xml)
                if match:
                    old = match.group(1)
                    new = '' if mutation == 'drawing-default' else 'draw'
                    xml = xml.replace(match.group(0), f'xmlns{":" + new if new else ""}="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"')
                    if old:
                        xml = re.sub(r'<(/?)' + re.escape(old) + ':', r'<\1' + (new + ':' if new else ''), xml)
                    elif new:
                        xml = re.sub(r'<(/?)([A-Za-z_][\w.-]*)(?=[\s/>])', r'<\1' + new + r':\2', xml)
                    files[path] = xml.encode()
    return fixed_zip(files)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    parser.add_argument('--seed', type=int, default=4863)
    parser.add_argument('--random-cases', type=int, default=300, help='Cases per independent writer')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    rng = random.Random(args.seed)
    manifest = {'seed': args.seed, 'writers': {'openpyxl': openpyxl.__version__, 'xlsxwriter': xlsxwriter.__version__}, 'cases': [], 'generatorErrors': []}
    base_values = [(0, 0, 'Item', 'value'), (0, 1, 'Amount', 'value'), (1, 0, 'A', 'value'),
                   (1, 1, 42, 'value'), (2, 0, 'B', 'value'), (2, 1, 15, 'value')]
    directed = [([], base_values, 'plain')]
    directed += [([feature], base_values, feature) for feature in FEATURES]
    directed += [(['chart', feature], base_values, 'chart-' + feature) for feature in ['comment', 'image', 'conditional']]
    directed += [([], base_values + [(4, 0, error, 'error')], 'error-' + str(index)) for index, error in enumerate(ERRORS)]
    directed += [([], base_values + [(4, 0, ('=1+2', value), 'formula')], 'cache-' + str(index)) for index, value in enumerate([3, 0, False, True, '', '文本', '#DIV/0!'])]
    directed += [([epoch] if epoch else [], base_values + [(4, 0, value, kind)], f'temporal-{epoch}-{index}') for epoch in ['', 'epoch1904'] for index, (value, kind) in enumerate([
        (datetime(2024, 2, 29, 12, 34, 56), 'date'), (time(12, 34, 56), 'time'), (timedelta(days=2, hours=3), 'time')])]
    directed += [([], base_values + [(4, 0, value, 'value')], 'text-' + str(index)) for index, value in enumerate(['<pic/><chart/> & 字符', '😀\n多行\t数据', '_x0041_', '_x000D_', '00123'])]

    def emit(data, cells, expected_features, writer, features, variant, valid=True):
        index = len(manifest['cases'])
        path = f'{index:05d}-{writer}-{variant}.xlsx'
        if valid:
            validate(data, cells)
        (args.output / path).write_bytes(data)
        manifest['cases'].append({'file': path, 'writer': writer, 'features': sorted(set(features)), 'variant': variant,
                                  'valid': valid, 'cells': cells, 'unsupportedFeatures': expected_features})

    for writer in ['openpyxl', 'xlsxwriter']:
        writer_directed = directed + ([(['textbox'], base_values, 'textbox')] if writer == 'xlsxwriter' else [])
        for features, values, label in writer_directed:
            try:
                data, cells, expected = workbook_bytes(writer, features, values, args.seed)
                emit(data, cells, expected, writer, features, label)
                if label in ['plain', 'chart', 'image', 'comment', 'conditional', 'chart-comment']:
                    for mutation in ['zip-store', 'zip-reverse-comment', 'sheet-prefix', 'sheet-filename', 'relationship-dot', 'relationship-absolute', 'drawing-default', 'drawing-alternate']:
                        try:
                            emit(mutate(data, mutation), cells, expected, writer, features, label + '--' + mutation)
                        except Exception as error:
                            manifest['generatorErrors'].append({'writer': writer, 'case': label, 'mutation': mutation, 'error': str(error)})
            except Exception as error:
                manifest['generatorErrors'].append({'writer': writer, 'case': label, 'error': str(error)})
        for index in range(args.random_cases):
            features = [feature for feature in FEATURES + (['textbox'] if writer == 'xlsxwriter' else []) if rng.random() < 0.3]
            values = copy.deepcopy(base_values)
            for row in range(5, rng.randint(6, 14)):
                value = rng.choice([rng.randint(-100000, 100000), rng.random() * 100, True, False, '中文😀', '<chart>literal</chart>', '00123'])
                values.append((row, rng.randrange(4), value, 'value'))
            values += [(3, 3, rng.choice(ERRORS), 'error')]
            try:
                data, cells, expected = workbook_bytes(writer, features, values, args.seed + index)
                mutation = rng.choice(['none', 'zip-store', 'zip-reverse-comment', 'drawing-default', 'drawing-alternate'])
                emit(mutate(data, mutation), cells, expected, writer, features, f'random-{index}-{mutation}')
            except Exception as error:
                manifest['generatorErrors'].append({'writer': writer, 'case': index, 'error': str(error)})

    plain = next(case for case in manifest['cases'] if case['variant'] == 'plain')
    source = (args.output / plain['file']).read_bytes()
    for index in range(128):
        data = bytearray(source)
        if index < 32:
            data = data[:rng.randrange(1, len(data))]
        else:
            for _ in range(rng.randint(1, 5)):
                data[rng.randrange(len(data))] ^= rng.randint(1, 255)
        try:
            validate(bytes(data), plain['cells'])
            valid = True
        except Exception:
            valid = False
        emit(bytes(data), plain['cells'], [], 'mutated', [], f'bytes-{index}', valid=valid)

    (args.output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'output': str(args.output), 'cases': len(manifest['cases']), 'valid': sum(case['valid'] for case in manifest['cases']), 'generatorErrors': len(manifest['generatorErrors'])}))


if __name__ == '__main__':
    main()
